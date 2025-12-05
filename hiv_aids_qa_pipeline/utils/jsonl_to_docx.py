import json
from docx import Document

def json_to_docx(json_path, docx_path):
    doc = Document()

    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    for i, item in enumerate(data):
        doc.add_heading(f"Page {i+1}", level=2)
        doc.add_paragraph(item)
        doc.add_paragraph("")

    doc.save(docx_path)
    print(f"Saved to {docx_path}")

json_to_docx("pages_clean.json", "output.docx")
